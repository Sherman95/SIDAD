import json
from datetime import date
from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook

from src.agent.critic import Critic
from src.agent.mock_llm import MockLLMProvider
from src.agent.orchestrator import Orchestrator
from src.agent.state import AgentState
from src.agent.tool_registry import get_default_registry
from src.models.powerbi_export import PowerBIExportResult


HEADERS = [
    "id_venta",
    "fecha",
    "producto",
    "categoria",
    "cantidad",
    "precio",
    "ciudad",
]
ROWS = [
    [1, date(2026, 1, 1), "Laptop", "Electronica", 1, 1200.5, "Quito"],
    [2, date(2026, 1, 2), "Mouse", None, 2, 25.0, "Cuenca"],
    [3, date(2026, 1, 3), "Silla", "Oficina", None, 150.0, "Loja"],
    [4, date(2026, 1, 4), "Escritorio", "Oficina", 1, 200.0, "Quito"],
    [4, date(2026, 1, 4), "Escritorio", "Oficina", 1, 200.0, "Quito"],
]


def _write_controlled_dataset(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix == ".csv":
        lines = [",".join(HEADERS)]
        for row in ROWS:
            lines.append(
                ",".join(
                    "" if value is None else value.isoformat() if isinstance(value, date) else str(value)
                    for value in row
                )
            )
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ventas"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    workbook.save(path)


def _run_agent(project_root: Path, suffix: str) -> tuple[AgentState, bytes]:
    raw_path = project_root / "data" / "raw" / f"ventas_controladas{suffix}"
    _write_controlled_dataset(raw_path)
    original_raw = raw_path.read_bytes()
    state = AgentState(
        user_request="Genera y exporta un dashboard Power BI integral.",
        input_dataset_path=f"data/raw/{raw_path.name}",
        current_goal="Validar el flujo completo de la Fase 12.5",
        project_base_path=str(project_root.resolve()),
        powerbi_output_path=f"data/output/phase-12-5-{suffix.removeprefix('.')}",
        max_iterations=15,
    )

    final_state = Orchestrator(
        llm=MockLLMProvider(),
        registry=get_default_registry(),
        critic=Critic(),
    ).run(state)
    return final_state, original_raw


def _assert_no_aabb_errors(state: AgentState) -> None:
    assert state.report_layout is not None
    assert state.validation_report is not None
    assert state.validation_report.is_valid
    assert not [issue for issue in state.validation_report.issues if issue.level == "Error"]

    for page in state.report_layout.pages:
        for index, visual in enumerate(page.placed_visuals):
            position = visual.position
            assert position.x + position.width <= page.canvas_width
            assert position.y + position.height <= page.canvas_height
            for other in page.placed_visuals[index + 1 :]:
                other_position = other.position
                separated = (
                    position.x >= other_position.x + other_position.width
                    or other_position.x >= position.x + position.width
                    or position.y >= other_position.y + other_position.height
                    or other_position.y >= position.y + position.height
                )
                assert separated


def _assert_export_artifacts(state: AgentState, project_root: Path) -> None:
    result = state.powerbi_export
    assert isinstance(result, PowerBIExportResult)
    assert result.success
    assert result.errors == []

    project_path = Path(result.project_path)
    semantic_path = Path(result.semantic_model_path)
    report_path = Path(result.report_path)
    assert project_path == project_root / project_path.relative_to(project_root)
    assert project_path.name == "SIDAD.pbip"
    assert project_path.is_file()
    assert semantic_path.name == "SIDAD.SemanticModel"
    assert report_path.name == "SIDAD.Report"

    semantic_definition = semantic_path / "definition"
    assert (semantic_definition / "database.tmdl").is_file()
    assert (semantic_definition / "model.tmdl").is_file()
    table_tmdl = semantic_definition / "tables" / "FACT_DATA.tmdl"
    assert table_tmdl.is_file()
    table_text = table_tmdl.read_text(encoding="utf-8")
    assert str(Path(state.processed_dataset_path).resolve()) in table_text

    report_definition = report_path / "definition"
    assert (report_path / "definition.pbir").is_file()
    assert (report_definition / "version.json").is_file()
    assert (report_definition / "report.json").is_file()
    assert (report_definition / "pages" / "pages.json").is_file()
    for json_path in [project_path, *report_path.rglob("*.json")]:
        json.loads(json_path.read_text(encoding="utf-8"))


def _report_identity(state: AgentState) -> tuple:
    pages_root = Path(state.powerbi_export.report_path) / "definition" / "pages"
    page_order = json.loads(
        (pages_root / "pages.json").read_text(encoding="utf-8")
    )["pageOrder"]
    visuals = tuple(
        (
            page_id,
            tuple(
                sorted(
                    path.parent.name
                    for path in (pages_root / page_id / "visuals").glob("*/visual.json")
                )
            ),
        )
        for page_id in page_order
    )
    return tuple(page_order), visuals


def test_csv_and_xlsx_run_end_to_end_with_deterministic_ids(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    unrelated_cwd = tmp_path / "unrelated-cwd"
    unrelated_cwd.mkdir()
    monkeypatch.chdir(unrelated_cwd)

    completed: dict[str, AgentState] = {}
    for suffix in (".csv", ".xlsx"):
        project_root = tmp_path / f"project-{suffix.removeprefix('.')}"
        project_root.mkdir()
        state, original_raw = _run_agent(project_root, suffix)
        completed[suffix] = state

        assert state.status == "SUCCESS"
        assert state.dataset_profile is not None
        assert state.dataset_profile.profile.row_count == 5
        assert state.dataset_profile.profile.duplicate_rows == 1
        assert state.processed_dataset_path is not None
        processed = pl.read_csv(state.processed_dataset_path)
        assert processed.height == 4
        assert processed.unique().height == 4
        assert processed["categoria"].null_count() == 0
        assert processed["cantidad"].null_count() == 0

        raw_path = project_root / "data" / "raw" / f"ventas_controladas{suffix}"
        assert raw_path.read_bytes() == original_raw
        assert Path(state.processed_dataset_path).is_relative_to(
            project_root / "data" / "processed"
        )

        assert state.semantic_project is not None
        assert state.semantic_project.tables[0].source_path == state.processed_dataset_path
        assert {kpi.name for kpi in state.semantic_project.kpis} == {
            "Total Registros",
            "Total Cantidad",
            "Total Precio",
            "Promedio de Precio",
        }
        assert state.dashboard_plan is not None
        assert [page.name for page in state.dashboard_plan.pages] == [
            "Resumen Ejecutivo",
            "Análisis Detallado",
        ]
        assert sum(
            len(page.visuals) + len(page.slicers)
            for page in state.dashboard_plan.pages
        ) == 12

        _assert_no_aabb_errors(state)
        _assert_export_artifacts(state, project_root.resolve())
        assert [record.tool_name for record in state.execution_history] == [
            "profile_dataset",
            "clean_dataset",
            "build_semantic_model",
            "create_dashboard_plan",
            "generate_layout",
            "validate_layout",
            "export_to_powerbi",
        ]
        assert state.execution_history[-1].is_error is False

    assert _report_identity(completed[".csv"]) == _report_identity(completed[".xlsx"])


@pytest.mark.parametrize("suffix", [".json", ".parquet", ".xls"])
def test_unsupported_formats_fail_explicitly(tmp_path: Path, suffix: str):
    source = tmp_path / f"dataset{suffix}"
    source.write_text("not a supported dataset", encoding="utf-8")

    is_error, result = get_default_registry().execute(
        "profile_dataset",
        {"file_path": str(source)},
    )

    assert is_error
    assert "Formato no soportado" in result
    assert "CSV" in result
    assert "XLSX" in result


def test_missing_dataset_fails_in_a_controlled_way(tmp_path: Path):
    missing = tmp_path / "missing.csv"

    is_error, result = get_default_registry().execute(
        "profile_dataset",
        {"file_path": str(missing)},
    )

    assert is_error
    assert "El archivo no existe" in result


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_empty_dataset_is_rejected_explicitly(tmp_path: Path, suffix: str):
    empty = tmp_path / f"empty{suffix}"
    if suffix == ".csv":
        empty.write_text("id,fecha,cantidad\n", encoding="utf-8")
    else:
        Workbook().save(empty)

    is_error, result = get_default_registry().execute(
        "profile_dataset",
        {"file_path": str(empty)},
    )

    assert is_error
    assert "El dataset esta vacio" in result


def test_cleaning_normalizes_xlsx_to_csv(tmp_path: Path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "processed" / "output.csv"
    _write_controlled_dataset(source)

    is_error, _ = get_default_registry().execute(
        "clean_dataset",
        {
            "input_path": str(source),
            "output_path": str(output),
            "actions": [{"action": "remove_duplicates"}],
        },
    )

    assert not is_error
    assert output.is_file()
    assert pl.read_csv(output).height == 4


def test_cleaning_rejects_non_csv_processed_output(tmp_path: Path):
    source = tmp_path / "input.csv"
    output = tmp_path / "processed" / "output.xlsx"
    _write_controlled_dataset(source)

    is_error, result = get_default_registry().execute(
        "clean_dataset",
        {"input_path": str(source), "output_path": str(output), "actions": []},
    )

    assert is_error
    assert "requiere un output_path CSV" in result
    assert not output.exists()
